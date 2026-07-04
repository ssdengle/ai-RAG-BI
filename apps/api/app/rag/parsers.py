from __future__ import annotations

from io import BytesIO
from pathlib import Path

from bs4 import BeautifulSoup

from apps.api.app.domain.documents import ParsedDocument
from apps.api.app.rag.exceptions import MissingParserDependencyError, UnsupportedDocumentTypeError
from apps.api.app.rag.interfaces import DocumentParser


class PlainTextParser:
    supported_extensions = (".txt",)
    supported_mime_types = ("text/plain",)

    def parse(self, filename: str, content_type: str | None, content: bytes) -> ParsedDocument:
        return ParsedDocument(text=_decode_text(content))


class MarkdownParser:
    supported_extensions = (".md", ".markdown")
    supported_mime_types = ("text/markdown", "text/x-markdown")

    def parse(self, filename: str, content_type: str | None, content: bytes) -> ParsedDocument:
        text = _decode_text(content)
        title = _extract_markdown_title(text)
        return ParsedDocument(text=text, title=title, extracted_metadata={"format": "markdown"})


class HtmlParser:
    supported_extensions = (".html", ".htm")
    supported_mime_types = ("text/html",)

    def parse(self, filename: str, content_type: str | None, content: bytes) -> ParsedDocument:
        soup = BeautifulSoup(content, "html.parser")

        for tag_name in ("script", "style", "noscript"):
            for tag in soup.find_all(tag_name):
                tag.decompose()

        body_text = soup.get_text(separator="\n", strip=True)
        title = soup.title.string.strip() if soup.title and soup.title.string else None
        description = None
        description_tag = soup.find("meta", attrs={"name": "description"})
        if description_tag is not None:
            description = description_tag.get("content")

        return ParsedDocument(
            text=body_text,
            title=title,
            extracted_metadata={
                "format": "html",
                "description": description,
            },
        )


class PdfParser:
    supported_extensions = (".pdf",)
    supported_mime_types = ("application/pdf",)

    def parse(self, filename: str, content_type: str | None, content: bytes) -> ParsedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - depends on local environment
            raise MissingParserDependencyError(
                "PDF ingestion requires the 'pypdf' package."
            ) from exc

        reader = PdfReader(BytesIO(content))
        pages: list[str] = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")

        title = None
        if reader.metadata is not None:
            title = getattr(reader.metadata, "title", None) or reader.metadata.get("/Title")

        return ParsedDocument(
            text="\n\n".join(pages),
            title=title,
            extracted_metadata={
                "format": "pdf",
                "page_count": len(reader.pages),
            },
        )


class DocxParser:
    supported_extensions = (".docx",)
    supported_mime_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    def parse(self, filename: str, content_type: str | None, content: bytes) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - depends on local environment
            raise MissingParserDependencyError(
                "DOCX ingestion requires the 'python-docx' package."
            ) from exc

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        core_properties = document.core_properties
        title = core_properties.title or None

        return ParsedDocument(
            text="\n\n".join(paragraphs),
            title=title,
            extracted_metadata={
                "format": "docx",
                "paragraph_count": len(paragraphs),
            },
        )


class DocumentParserRegistry:
    def __init__(self, parsers: list[DocumentParser]) -> None:
        self._parsers = parsers

    def resolve(self, filename: str, content_type: str | None) -> DocumentParser:
        extension = Path(filename).suffix.lower()
        normalized_content_type = (content_type or "").split(";")[0].strip().lower()

        for parser in self._parsers:
            if extension in parser.supported_extensions:
                return parser

        for parser in self._parsers:
            if normalized_content_type and normalized_content_type in parser.supported_mime_types:
                return parser

        raise UnsupportedDocumentTypeError(
            f"Unsupported document type for filename '{filename}' and content type '{content_type}'."
        )


def build_default_parser_registry() -> DocumentParserRegistry:
    return DocumentParserRegistry(
        parsers=[
            PlainTextParser(),
            MarkdownParser(),
            HtmlParser(),
            PdfParser(),
            DocxParser(),
        ]
    )


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue

    return content.decode("utf-8", errors="replace")


def _extract_markdown_title(text: str) -> str | None:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            return stripped.lstrip("#").strip() or None
    return None
